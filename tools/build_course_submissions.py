"""Build reviewable Moodle folders and native PPTX/PDF presentations.

Optional maintainer utility only; the delivered site/prototype needs no Python,
package manager, build or server. Source narratives live in course-source/.
Run without --include-web while the current website is still being edited.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import shutil
import xml.etree.ElementTree as ET
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'course-source'
WIDTH, HEIGHT = 960, 540
INK, PAPER, ORANGE, MUTED = '#f5f6fa', '#101217', '#ff885d', '#b7bfcc'
PALE, LINE, MINT = '#1d222c', '#3d4552', '#ade6d5'
FONT = 'Arial'
# Liberation Sans is metrically compatible with Arial. Both exports therefore use
# the same explicit line breaks; the PDF embeds a font with all Czech glyphs.
FONT_DIR = ROOT / 'tools/fonts'
AUTHORS = ['Michal Ladra', 'Matěj Mlejnek', 'Jan Martínek']


def rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color.removeprefix('#'))


def data_checks() -> tuple[dict, list[str]]:
    text = (ROOT / 'web/js/data.js').read_text(encoding='utf-8')
    data = json.loads(text[text.index('{'):].rstrip().removesuffix(';'))
    xml = ET.parse(ROOT / 'data.xml').getroot()
    checks = []
    for name, xpath, minimum in [('dny', './dny/den', 3), ('scény', './haly/hala', 3),
                                ('interpreti', './interpreti/interpret', 60),
                                ('vystoupení', './sumakce/akce', 60)]:
        nodes = xml.findall(xpath)
        assert len(nodes) >= minimum, f'Missing minimum {name}'
        assert len({node.attrib['id'] for node in nodes}) == len(nodes), f'Duplicate IDs {name}'
        checks.append(f'{len(nodes)} {name}; jedinečná ID')
    ids = {kind: {node.attrib['id'] for node in xml.findall(path)} for kind, path in
           [('day', './dny/den'), ('venue', './haly/hala'), ('artist', './interpreti/interpret')]}
    for event in xml.findall('./sumakce/akce'):
        assert event.findtext('den_ref') in ids['day']
        assert event.findtext('hala_ref') in ids['venue']
        assert all(node.text in ids['artist'] for node in event.findall('./interpreti/interpret_ref'))
    checks.append('Všechny XML reference dne, scény a interpreta existují.')
    assert data['festival']['name'] == xml.findtext('nazev')
    assert len(data['performers']) == len(ids['artist'])
    assert len(data['events']) == len(xml.findall('./sumakce/akce'))
    assert {item['id'] for item in data['venues']} == ids['venue']
    for event in data['events']:
        assert event['dayId'] in ids['day'] and event['venueId'] in ids['venue']
        assert all(item in ids['artist'] for item in event['performerIds'])
    checks.append('Přibalená data mají shodný název festivalu, počty a platné reference.')
    return data, checks


class Deck:
    """Render one festival identity to editable PowerPoint and matching PDF.

    Slide text stays short; the explanation belongs in presenter notes. Layout
    measurement uses the actual embedded PDF font rather than a character count.
    The builder fails when text will not fit at a readable size, so editing the
    narrative cannot silently create clipped slides.
    """
    def __init__(self, target: Path, number: str, info: dict):
        self.target, self.number, self.info = target, number, info
        self.ppt = Presentation()
        self.ppt.slide_width, self.ppt.slide_height = Pt(WIDTH), Pt(HEIGHT)
        self.ppt.core_properties.title = f"Prague Pulse Fest — cvičení {number}: {info['name']}"
        self.ppt.core_properties.subject = 'Webdesign: návrh, data, uživatelské scénáře a ověření Prague Pulse Fest.'
        self.authors = info.get('authors', AUTHORS)
        self.ppt.core_properties.author = ', '.join(self.authors)
        pdfmetrics.registerFont(TTFont('PRF-Regular', str(FONT_DIR / 'LiberationSans-Regular.ttf')))
        pdfmetrics.registerFont(TTFont('PRF-Bold', str(FONT_DIR / 'LiberationSans-Bold.ttf')))
        self.pdf = canvas.Canvas(str(target / 'prezentace.pdf'), pagesize=(WIDTH, HEIGHT))
        self.pdf.setTitle(self.ppt.core_properties.title)
        self.pdf.setAuthor(', '.join(self.authors))
        self.slide = None
        self.notes = []
        self.overflow = []

    def rect(self, x, y, w, h, color, border=None):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color)
        if border:
            shape.line.color.rgb = rgb(border); shape.line.width = Pt(1.1)
        else:
            shape.line.fill.background()
        self.pdf.setFillColor(color)
        if border:
            self.pdf.setStrokeColor(border); self.pdf.setLineWidth(1.1)
        self.pdf.rect(x, HEIGHT-y-h, w, h, fill=1, stroke=bool(border))

    def lines(self, text, w, size, bold=False):
        """Wrap words by measured width, including long filenames as a fallback."""
        font = 'PRF-Bold' if bold else 'PRF-Regular'
        lines = []
        # A small reserve also tolerates font substitution when a PPTX is opened
        # on a computer without Arial installed.
        usable = w * .98
        for source in str(text).split('\n'):
            if not source:
                lines.append(''); continue
            line = ''
            for word in source.split():
                candidate = f'{line} {word}' if line else word
                if pdfmetrics.stringWidth(candidate, font, size) <= usable:
                    line = candidate; continue
                if line:
                    lines.append(line); line = ''
                if pdfmetrics.stringWidth(word, font, size) > usable:
                    for char in word:
                        if pdfmetrics.stringWidth(line + char, font, size) > usable:
                            lines.append(line); line = ''
                        line += char
                else:
                    line = word
            if line:
                lines.append(line)
        return lines or ['']

    def text(self, x, y, w, text, size=24, color=INK, bold=False, h=None, spacing=1.18):
        font = 'PRF-Bold' if bold else 'PRF-Regular'
        lines = self.lines(text, w, size, bold)
        needed = size + max(0, len(lines)-1) * size * spacing + 5
        height = needed if h is None else h
        if (needed > height + .1 or min(x,y) < 0 or x+w > WIDTH+.1
                or y+height > HEIGHT+.1):
            self.overflow.append((self.number, str(text)[:60], round(needed, 1), height))
        shape = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(height))
        frame = shape.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = FONT
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = rgb(color)
            paragraph.space_after = Pt(0)
            paragraph.line_spacing = Pt(size * spacing)
        self.pdf.setFont(font, size); self.pdf.setFillColor(color)
        ascent = pdfmetrics.getAscent(font) / 1000 * size
        for index, line in enumerate(lines):
            self.pdf.drawString(x, HEIGHT-y-ascent-index*size*spacing, line)
        return needed

    def fit(self, text, w, h, size=28, minimum=24, bold=False, spacing=1.18):
        while size >= minimum:
            count = len(self.lines(text, w, size, bold))
            if size + max(0,count-1)*size*spacing + 5 <= h:
                return size
            size -= .5
        raise ValueError(f'Cvičení {self.number}: zkraťte text snímku: {str(text)[:100]}')

    def connector(self, x1, y1, x2, y2, color=LINE):
        """The same vector connector is written to both native output formats."""
        shape = self.slide.shapes.add_connector(1, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        shape.line.color.rgb = rgb(color)
        shape.line.width = Pt(1.5)
        self.pdf.setStrokeColor(color); self.pdf.setLineWidth(1.5)
        self.pdf.line(x1, HEIGHT-y1, x2, HEIGHT-y2)

    def footer(self, index):
        self.rect(48,494,864,1,LINE)
        label = ('ŠABLONA · doplnit skutečný cizí projekt' if self.number == '06'
                 else 'Prague Pulse Fest · Výstaviště Praha · 2027')
        self.text(48,509,780,label,11,MUTED)
        self.text(864,507,48,f'{index} / 6',12,MUTED)

    def bullets(self, values, top, bottom, lead=''):
        if not 1 <= len(values) <= 4:
            raise ValueError(f'Cvičení {self.number}: použijte 1–4 body na snímek.')
        lead_h = 0
        if lead:
            lead_h = 49 if len(self.lines(lead,848,18)) > 1 else 28
        available = bottom-top-lead_h-(18 if lead else 0)
        gap = 20
        size = 28
        while size >= 24:
            heights = [size + max(0,len(self.lines(v,800,size))-1)*size*1.2+5
                       for v in values]
            if sum(heights)+gap*(len(values)-1) <= available:
                break
            size -= .5
        if size < 24:
            raise ValueError(f'Cvičení {self.number}: body jsou příliš dlouhé pro čitelný snímek.')
        y = top
        for value,height in zip(values,heights):
            self.rect(56,y+10,8,8,ORANGE)
            self.text(84,y,800,value,size,INK,h=height,spacing=1.2)
            y += height+gap
        if lead:
            self.text(56,bottom-lead_h,848,lead,18,MUTED,h=lead_h)

    def audience(self, item, top):
        for i,(heading,question,answer) in enumerate(item['cards']):
            x=48+i*294
            self.rect(x,top,276,249,PALE)
            self.rect(x,top,276,4,ORANGE if i!=1 else MINT)
            self.text(x+18,top+18,240,f'0{i+1}',15,ORANGE,True)
            self.text(x+18,top+50,240,heading,24,INK,True,h=61)
            self.text(x+18,top+121,240,question,20,MINT,h=56)
            self.text(x+18,top+184,240,answer,19,MUTED,h=56)
        self.text(48,top+267,864,item['lead'],18,MUTED,h=38)

    def data_venues(self, item, top):
        # Show the same ID relationship that drives filters and ticket prefills.
        for i,(value,label) in enumerate(item['metrics']):
            x=48+i*221
            self.text(x,top-3,200,value,45,ORANGE,True,h=55)
            self.text(x,top+46,200,label,18,INK,h=28)
        for i,(heading,body) in enumerate(item['entities']):
            x=48+i*221
            if i:
                self.connector(x-14,top+130,x-4,top+130,ORANGE)
            self.rect(x,top+93,201,99,PALE)
            self.rect(x,top+93,201,3,MINT if i==1 else ORANGE)
            self.text(x+13,top+108,175,heading,23,INK,True,h=32)
            self.text(x+13,top+149,175,body,16,MUTED,h=35)
        self.text(48,top+219,864,item['pipeline'],20,MINT,True,h=32)
        self.text(48,top+262,864,item['lead'],17,MUTED,h=39)

    def sitemap(self, item, top):
        # Branching lines show a sitemap; content links remain explicit in cards.
        center=480; root_y=top-7
        self.rect(352,root_y,256,41,ORANGE)
        self.text(370,root_y+7,220,item['root'],23,PAPER,True,h=29)
        bar_y=root_y+54
        self.connector(center,root_y+41,center,bar_y)
        # Six siblings share root branches. The vertical trunks stay outside
        # cards, so the second row cannot appear to be a child of the first.
        self.connector(38,bar_y,626,bar_y)
        for col in range(3):
            self.connector(38+294*col,bar_y,38+294*col,root_y+69+86+38)
        for i,(heading,body) in enumerate(item['nodes']):
            col,row=i%3,i//3; x=48+294*col; y=root_y+69+row*86
            self.connector(x-10,y+38,x,y+38)
            self.rect(x,y,276,77,PALE)
            self.text(x+15,y+10,246,heading,24,INK,True,h=34)
            self.text(x+15,y+48,246,body,16,MUTED,h=25)
        self.text(48,top+259,864,item['lead'],18,MUTED,h=38)

    def scenarios(self, item, top):
        for i,scenario in enumerate(item['scenarios']):
            x=48+i*442
            self.rect(x,top,422,256,PALE)
            self.rect(x,top,52,39,ORANGE if i==0 else MINT)
            self.text(x+10,top+7,38,scenario['id'],21,PAPER,True,h=29)
            self.text(x+66,top+8,339,scenario['title'],23,INK,True,h=34)
            self.text(x+18,top+63,386,scenario['who'],20,MINT,True,h=29)
            self.text(x+18,top+101,386,scenario['start'],19,MUTED,h=32)
            self.text(x+18,top+148,386,scenario['steps'],21,INK,h=58)
            self.rect(x+18,top+219,386,1,LINE)
            self.text(x+18,top+231,386,scenario['result'],18,ORANGE,h=25)
        self.text(48,top+274,864,item['lead'],18,MUTED,h=33)

    def verification(self, item, top):
        for i,(label,body) in enumerate(item['decisions']):
            x=48+i*294
            self.rect(x,top,276,113,PALE)
            self.text(x+18,top+16,240,label,23,ORANGE,True,h=33)
            self.text(x+18,top+59,240,body,20,INK,h=49)
        self.text(48,top+145,864,'Konkrétní ověření dat',22,MINT,True,h=32)
        for i,label in enumerate(item['checks']):
            x=48+i*294
            self.text(x,top+193,31,f'{i+1}.',21,ORANGE,True,h=31)
            self.text(x+34,top+193,242,label,21,INK,h=56)
        self.text(48,top+269,864,item['lead'],18,MUTED,h=33)

    def page(self, item, index):
        self.slide = self.ppt.slides.add_slide(self.ppt.slide_layouts[6])
        cover = item['layout']=='cover'
        self.rect(0,0,WIDTH,HEIGHT,PAPER)
        self.notes.append(item['notes'])
        self.slide.notes_slide.notes_text_frame.text = item['notes']
        if cover:
            # An original type-based festival cover, fully editable in PowerPoint.
            self.rect(0,0,12,HEIGHT,ORANGE)
            self.rect(650,0,310,HEIGHT,PALE)
            self.rect(650,0,310,8,ORANGE)
            self.text(48,43,864,item['kicker'],14,ORANGE,True,h=24)
            self.text(45,94,865,'PRAGUE\nPULSE FEST',66,INK,True,h=148,spacing=.98)
            size=self.fit(item['title'],855,86,38,31,True,1.06)
            self.text(48,275,855,item['title'],size,ORANGE,True,h=86,spacing=1.06)
            self.text(48,385,855,item['lead'],23,MUTED,h=65)
            self.text(48,463,864,' · '.join(self.authors),18,INK,h=26)
        else:
            self.rect(0,0,8,HEIGHT,ORANGE)
            self.text(48,35,864,item['kicker'],13,MUTED,True,h=22)
            self.rect(48,69,54,4,ORANGE)
            title_size=self.fit(item['title'],864,89,36,30,True,1.12)
            self.text(46,95,866,item['title'],title_size,INK,True,h=89,spacing=1.12)
            title_lines=len(self.lines(item['title'],864,title_size,True))
            top=174 if title_lines==1 else 204
            layout = item['layout']
            if layout == 'audience': self.audience(item,top)
            elif layout == 'data-venues': self.data_venues(item,top)
            elif layout == 'sitemap': self.sitemap(item,top)
            elif layout == 'scenarios': self.scenarios(item,top)
            elif layout == 'verification': self.verification(item,top)
            elif layout == 'bullets':
                self.bullets(item['bullets'],top,474,item.get('lead',''))
            elif layout in ('cards','split'):
                values=[f'{heading}: {body.replace(chr(10), " ")}'
                        for heading,body in item['cards']]
                self.bullets(values,top,474,item.get('lead',''))
            elif layout=='metrics':
                assert len(item['metrics']) == 4, 'Metrics layout expects four values'
                for i,(value,label) in enumerate(item['metrics']):
                    x=56+i*218
                    self.text(x,top+20,191,value,72,ORANGE,True,h=90)
                    self.text(x,top+123,191,label,22,INK,h=68)
                self.text(56,407,848,item.get('lead',''),22,MUTED,h=65)
            elif layout=='route':
                assert len(item['steps']) <= 4, 'Route layout supports at most four steps'
                for i,(heading,body) in enumerate(item['steps']):
                    y=top+i*(65 if title_lines==1 else 57)
                    self.rect(48,y,48,38,PALE)
                    self.text(58,y+6,38,f'{i+1:02}',21,ORANGE,True,h=29)
                    self.text(112,y+3,268,heading,24,INK,True,h=54)
                    self.text(391,y+5,513,body,23,MUTED,h=55)
                if item.get('lead'):
                    self.text(48,445,864,item['lead'],18,MUTED,h=40)
            else:
                raise ValueError(f'Unknown slide layout: {layout}')
        self.footer(index)
        self.pdf.showPage()

    def finish(self):
        assert not self.overflow, f'Presentation overflow: {self.overflow}'
        self.ppt.save(self.target / 'prezentace.pptx')
        self.pdf.save()
        notes = [f"# Cvičení {self.number} — poznámky k prezentaci",'',
                 self.info['name'],'',
                 '**Autoři:** '+', '.join(self.authors),'',
                 'Šest snímků, přibližně 35–45 sekund na snímek; celkem 3–5 minut.',
                 'Poznámky jsou příprava k mluvenému projevu; na Moodle se neodevzdávají.','']
        if self.number=='06':
            notes += ['**Nejde o dokončené peer review. Nejdříve doplňte cizí projekt a skutečné nálezy.**','']
        for i,(slide,note) in enumerate(zip(self.info['slides'],self.notes),1):
            notes += [f"## {i}. {slide['title'].replace(chr(10),' ')}",'',note,'']
        (self.target / 'presenter-notes.md').write_text('\n'.join(notes),encoding='utf-8')


def copy_file(source: Path, target: Path):
    target.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(source,target)


def build(output: Path, include_web: bool):
    data, checks = data_checks()
    output = output.resolve()
    assert output != ROOT and output not in ROOT.parents, 'Choose a separate output folder'
    marker = output / '.course-generated'
    if output.exists() and any(output.iterdir()) and not marker.exists():
        raise SystemExit(f'Output {output} is not an owned generated folder; choose an empty directory.')
    output.mkdir(parents=True,exist_ok=True)
    marker.write_text('Generated from course-source by tools/build_course_submissions.py\n')
    folder_names = {'02':'cviceni-02-specifikace','03':'cviceni-03-prototyp',
                    '04':'cviceni-04-implementace','05':'cviceni-05-review-ready',
                    '06':'cviceni-06-cross-review'}
    deck_source = json.loads((SOURCE / 'presentations.json').read_text(encoding='utf-8'))
    for number,name in folder_names.items():
        folder=output/name
        if folder.exists():shutil.rmtree(folder)
        folder.mkdir()
        info=deck_source[number]
        deck=Deck(folder,number,info)
        for i,item in enumerate(info['slides'],1):deck.page(item,i)
        deck.finish()
        if number == '02':
            # Keep the upload directory minimal; previews and preparation stay separate.
            for subdir, source_name, output_name in [
                ('nahledy','prezentace.pdf','cviceni-02.pdf'),
                ('poznamky','presenter-notes.md','cviceni-02.md')]:
                extra=output/subdir; extra.mkdir(exist_ok=True)
                shutil.move(str(folder/source_name),str(extra/output_name))
    c02=output/folder_names['02']; c03=output/folder_names['03']
    c04=output/folder_names['04']; c05=output/folder_names['05']; c06=output/folder_names['06']
    for folder in [c02,c05]:
        copy_file(ROOT/'project-specification.md',folder/'project-specification.md')
        copy_file(ROOT/'data.xml',folder/'data.xml')
        if (ROOT/'data/festival.xsd').exists():copy_file(ROOT/'data/festival.xsd',folder/'festival.xsd')
    # The first submission needs one presentation, its data and specification.
    # Keep speaker notes and duplicate PDF exports out of this upload package.
    with zipfile.ZipFile(output/'cviceni-02-k-odevzdani.zip','w',zipfile.ZIP_DEFLATED) as archive:
        for name in ['project-specification.md','data.xml','festival.xsd','prezentace.pptx']:
            archive.write(c02/name,arcname=name)
    for source,target in [('c03-progress-02.md',c03/'progress-02.md'),
                          ('c04-progress-03.md',c04/'progress-03.md'),
                          ('c05-progress-04.md',c05/'progress-04.md'),
                          ('c06-review-template.md',c06/'review.md')]:
        copy_file(SOURCE/source,target)
    prototype=c03/'prototype'
    shutil.copytree(SOURCE/'prototype',prototype,dirs_exist_ok=True)
    copy_file(ROOT/'web/js/data.js',prototype/'data.js')
    if (SOURCE/'prototype-validation.md').exists():
        copy_file(SOURCE/'prototype-validation.md',c03/'prototype-validation.md')
    if include_web:
        for folder in [c04,c05]:
            shutil.copytree(ROOT/'web',folder/'web')
            for name in ['documentation.md','validation_report.md']:
                if (ROOT/name).exists():copy_file(ROOT/name,folder/name)
            if (ROOT/'preview/home.png').exists():
                copy_file(ROOT/'preview/home.png',folder/'preview/home.png')
        final=output/'finalni-projekt'
        if final.exists():shutil.rmtree(final)
        project=final/'projekt';project.mkdir(parents=True)
        for name in ['project-specification.md','documentation.md','data.xml','validation_report.md']:
            if (ROOT/name).exists():copy_file(ROOT/name,project/name)
        for name in ['documentation.docx','documentation.pdf']:
            if (ROOT/'docs'/name).exists():copy_file(ROOT/'docs'/name,project/'docs'/name)
        if (ROOT/'preview/home.png').exists():
            copy_file(ROOT/'preview/home.png',project/'preview/home.png')
        if (ROOT/'data/festival.xsd').exists():copy_file(ROOT/'data/festival.xsd',project/'festival.xsd')
        shutil.copytree(ROOT/'web',project/'web')
        shutil.copytree(prototype,project/'prototype')
        if (ROOT/'review/received-review.md').exists():copy_file(ROOT/'review/received-review.md',project/'review/received-review.md')
        else:
            (project/'review').mkdir()
            (project/'review/received-review.md').write_text('# Přijaté cross-review\n\nSkutečné hodnocení jiným týmem nebylo poskytnuto. Doplnit po provedení.\n',encoding='utf-8')
        review_names = ['internal-review.md', 'audio-requirements-2026-10-18.md']
        review_names += [path.name for path in sorted((ROOT/'review').glob('*-validation-*.json'))]
        for review_name in review_names:
            if (ROOT/'review'/review_name).exists():copy_file(ROOT/'review'/review_name,project/'review'/review_name)
        (final/'README.md').write_text('# Finální projekt\n\nSložka `projekt/` odpovídá povinné struktuře zadání. Otevřete `projekt/web/index.html`.\n\nPřed odevzdáním doplňte skutečné přijaté review, týmovou historii a potvrzení rozsahu změn.\n',encoding='utf-8')
    (output/'data-validation.md').write_text('# Kontrola dat při sestavení podkladů\n\n'+
        '\n'.join('- '+item for item in checks)+'\n\nKontrola se spouští při každém sestavení. Nejde o HTML validaci ani UX audit.\n',encoding='utf-8')
    readme = """# Podklady k odevzdání

## Cvičení 2

Použijte **`cviceni-02-k-odevzdani.zip`**, nebo nahrajte obsah
`cviceni-02-specifikace/`. Obsahuje pouze čtyři soubory:

- `project-specification.md`
- `data.xml`
- `festival.xsd`
- `prezentace.pptx`

Poznámky jsou zvlášť v `poznamky/cviceni-02.md` a také uvnitř PowerPointu.
PDF náhled je v `nahledy/cviceni-02.pdf`. Tyto soubory k odevzdání nepřidávejte,
pokud PDF nepoužijete místo PowerPointu k promítání.

## Další výstupy

| Složka | Obsah k odevzdání |
| --- | --- |
| `cviceni-03-prototyp/` | `prototype/`, `progress-02.md`, použitá prezentace |
| `cviceni-04-implementace/` | `web/`, `progress-03.md`, použitá prezentace |
| `cviceni-05-review-ready/` | `web/`, specifikace, `progress-04.md`, použitá prezentace |
| `cviceni-06-cross-review/` | `review.md` a prezentace nálezů cizího projektu; zatím šablona |
| `finalni-projekt/projekt/` | Specifikace, data, prototyp, web, dokumentace a přijaté review |

Prezentace mají šest snímků a poznámky na 3–5 minut. Pro cvičení 3–6 je
přiložen PowerPoint i PDF; odevzdejte formát, který budete promítat.
Složky webu a prototypu vždy ponechte celé.

`manifest.json` zaznamenává soubory a jejich kontrolní otisky. Záznamy
`data-validation.md` a `prototype-validation.md` slouží ke kontrole podkladů.
Cvičení 6 vyžaduje doplnit skutečné nálezy přiděleného projektu; vlastní
kontroly se nevydávají za přijaté externí review.
"""
    if not include_web:
        readme += '\n**Pracovní export:** současný web nebyl kopírován. Úplný export vyžaduje `--include-web`.\n'
    (output/'README.md').write_text(readme,encoding='utf-8')
    if (SOURCE/'prototype-validation.md').exists():
        copy_file(SOURCE/'prototype-validation.md',output/'prototype-validation.md')
    files=[]
    for path in sorted(output.rglob('*')):
        if path.is_file() and path.name not in {'.course-generated','manifest.json'}:
            files.append({'path':str(path.relative_to(output)), 'bytes':path.stat().st_size,
                          'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    (output/'manifest.json').write_text(json.dumps({'complete_web_copies':include_web,'files':files},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f'Built {len(files)} files in {output}; website copies: {include_web}')
    return output


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=ROOT/'odevzdani')
    parser.add_argument('--include-web',action='store_true')
    args=parser.parse_args()
    build(args.output,args.include_web)
