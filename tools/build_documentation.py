"""Create DOCX/PDF companions from the required documentation.md.

Optional authoring tool: running or reading the delivered website never needs it.
The Markdown remains the editable source; rerun after changing its content.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from html import escape
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
AUTHORS = 'Michal Ladra, Matěj Mlejnek, Jan Martínek'
TITLE = 'Prague Pulse Fest'
ORANGE = 'E4552E'
INK = '242424'
MUTED = '666666'
# Keep the cover aligned with the supplied earlier documentation; do not carry
# forward its old authors, date, academic year or course name.
INSTITUTION = ['Česká zemědělská univerzita v Praze', 'Provozně ekonomická fakulta',
               'Katedra informačních technologií']
INLINE = re.compile(r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)')


def blocks(source: str):
    """Parse only the small Markdown subset actually used by documentation.md."""
    lines = source.splitlines()
    i = next(i for i, line in enumerate(lines) if line.startswith('## '))
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
        elif line.startswith('```'):
            end = i + 1
            while end < len(lines) and not lines[end].startswith('```'):
                end += 1
            yield 'code', '\n'.join(lines[i + 1:end])
            i = end + 1
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [cell.strip() for cell in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in cells):
                    rows.append(cells)
                i += 1
            yield 'table', rows
        elif line.startswith('### '):
            yield 'h3', line[4:]
            i += 1
        elif line.startswith('## '):
            yield 'h2', line[3:]
            i += 1
        elif line.startswith('!['):
            match = re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)', line)
            if match:
                yield 'image', (match[1], ROOT / match[2])
            i += 1
        elif re.match(r'^(?:- |\d+\. )', line):
            ordered = not line.startswith('- ')
            item = re.sub(r'^(?:- |\d+\. )', '', line)
            i += 1
            while i < len(lines) and lines[i].startswith('  '):
                item += ' ' + lines[i].strip()
                i += 1
            yield 'number' if ordered else 'bullet', item
        else:
            paragraph = [line.strip()]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r'^(?:#|\||!\[|```|- |\d+\. )', lines[i]):
                paragraph.append(lines[i].strip())
                i += 1
            yield 'p', ' '.join(paragraph)


def add_word_link(paragraph, label, url):
    part = paragraph.part
    relation = part.relate_to(url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), relation)
    run = OxmlElement('w:r')
    prop = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), ORANGE); prop.append(color)
    run.append(prop)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); paragraph._p.append(link)


def word_inline(paragraph, text, size=None):
    for token in INLINE.split(text):
        match = re.fullmatch(r'\[([^]]+)\]\(([^)]+)\)', token)
        if match:
            add_word_link(paragraph, match[1], match[2])
            continue
        run = paragraph.add_run(token[2:-2] if token.startswith('**') else token[1:-1] if token.startswith('`') else token)
        if token.startswith('**'):
            run.bold = True
        if token.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(size or 9)
        elif size:
            run.font.size = Pt(size)


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), color); props.append(fill)


def render_docx(items, output):
    doc = Document()
    props = doc.core_properties
    props.title = TITLE + ' — dokumentace projektu Webdesign'
    props.author = AUTHORS
    props.subject = 'Akademický rok 2026/2027; klientský festivalový web'
    props.created = datetime(2026, 10, 18, tzinfo=timezone.utc)
    props.modified = datetime(2026, 10, 19, tzinfo=timezone.utc)
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(2); section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.1); section.right_margin = Cm(2.1)
    section.different_first_page_header_footer = True
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'; normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)
    for style_name, font_size in [('Title', 36), ('Heading 1', 19), ('Heading 2', 13)]:
        style = doc.styles[style_name]
        style.font.name = 'Arial'; style.font.size = Pt(font_size)
        style.font.color.rgb = RGBColor.from_string(ORANGE if style_name == 'Heading 1' else INK)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(8)
    header = section.header.paragraphs[0]
    header.text = 'PRAGUE PULSE FEST  /  WEBDESIGN'
    header.style = doc.styles['Caption']
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Dokumentace projektu · ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    for line in INSTITUTION:
        paragraph = doc.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
    doc.add_paragraph().paragraph_format.space_after = Pt(82)
    paragraph = doc.add_paragraph('Dokumentace projektu', 'Subtitle')
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = doc.add_paragraph(TITLE, 'Title'); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = doc.add_paragraph('Webdesign · akademický rok 2026/2027')
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = doc.add_paragraph('Fiktivní festivalový web\nHTML · CSS · JavaScript · XML')
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(70)
    for line in [AUTHORS, '19. října 2026', 'Festival: 1.–3. srpna 2027']:
        paragraph = doc.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    doc.add_heading('Obsah', level=1)
    for kind, value in items:
        if kind == 'h2':
            doc.add_paragraph(value)
    doc.add_paragraph('Povinným zdrojovým dokumentem je documentation.md. Tato verze zachovává stejný obsah v upravitelném formátu pro čtení a odevzdání.')
    doc.add_page_break()
    for kind, value in items:
        if kind in ('h2', 'h3'):
            doc.add_heading(value, level=1 if kind == 'h2' else 2)
        elif kind == 'table':
            table = doc.add_table(rows=0, cols=len(value[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            width = 16.8 / len(value[0])
            for column in table.columns:
                column.width = Cm(width)
            for index, row in enumerate(value):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.width = Cm(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.space_before = Pt(4)
                    paragraph.paragraph_format.line_spacing = 1.1
                    word_inline(paragraph, text, size=8.4)
                    if index == 0:
                        shade(cell, 'F3DFD7')
                        for run in paragraph.runs:
                            run.bold = True
                    elif index % 2 == 0:
                        shade(cell, 'F5F5F5')
                props = table.rows[-1]._tr.get_or_add_trPr()
                no_split = OxmlElement('w:cantSplit'); props.append(no_split)
                if index == 0:
                    repeat = OxmlElement('w:tblHeader'); props.append(repeat)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif kind == 'image':
            caption, path = value
            if path.exists():
                doc.add_picture(str(path), width=Cm(16.4))
                paragraph = doc.add_paragraph(caption, 'Caption')
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == 'code':
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run(value); run.font.name = 'Consolas'; run.font.size = Pt(9)
        else:
            paragraph = doc.add_paragraph(style='List Bullet' if kind == 'bullet' else 'List Number' if kind == 'number' else 'Normal')
            word_inline(paragraph, value)
    doc.save(output)


def pdf_inline(text):
    parts = []
    for token in INLINE.split(text):
        match = re.fullmatch(r'\[([^]]+)\]\(([^)]+)\)', token)
        if match:
            parts.append(f'<link href="{escape(match[2], quote=True)}" color="#{ORANGE}">{escape(match[1])}</link>')
        elif token.startswith('**'):
            parts.append('<b>' + escape(token[2:-2]) + '</b>')
        elif token.startswith('`'):
            parts.append('<font size="9">' + escape(token[1:-1]) + '</font>')
        else:
            parts.append(escape(token))
    return ''.join(parts)


def render_pdf(items, output):
    fonts = ROOT / 'tools/fonts'
    pdfmetrics.registerFont(TTFont('PulseDoc', str(fonts / 'LiberationSans-Regular.ttf')))
    pdfmetrics.registerFont(TTFont('PulseDoc-Bold', str(fonts / 'LiberationSans-Bold.ttf')))
    pdfmetrics.registerFontFamily('PulseDoc', normal='PulseDoc', bold='PulseDoc-Bold')
    style = ParagraphStyle('body', fontName='PulseDoc', fontSize=10, leading=14,
                           textColor=colors.HexColor('#'+INK), spaceAfter=8)
    h2 = ParagraphStyle('chapter', parent=style, fontName='PulseDoc-Bold', fontSize=18,
                        leading=22, textColor=colors.HexColor('#'+ORANGE), spaceBefore=17,
                        spaceAfter=10, keepWithNext=True)
    h3 = ParagraphStyle('subchapter', parent=h2, fontSize=12, leading=16,
                        textColor=colors.HexColor('#'+INK), spaceBefore=12, spaceAfter=7)
    cell = ParagraphStyle('cell', parent=style, fontSize=8.2, leading=11, spaceAfter=0,
                          wordWrap='CJK')
    centered = ParagraphStyle('center', parent=style, alignment=TA_CENTER, spaceAfter=5)
    caption_style = ParagraphStyle('caption', parent=centered, fontSize=8, leading=11,
                                   textColor=colors.HexColor('#'+MUTED))
    width = 463.28  # A4 content width, including Platypus frame padding
    story = [Spacer(1, 28)]
    for line in INSTITUTION:
        story.append(Paragraph(escape(line), centered))
    story.extend([Spacer(1, 95), Paragraph('Dokumentace projektu', centered),
                  Paragraph(TITLE, ParagraphStyle('cover', parent=centered, fontName='PulseDoc-Bold', fontSize=35, leading=43)),
                  Spacer(1, 16), Paragraph('Webdesign · akademický rok 2026/2027', centered),
                  Paragraph('Fiktivní festivalový web', centered),
                  Paragraph('HTML · CSS · JavaScript · XML', centered), Spacer(1, 90),
                  Paragraph(escape(AUTHORS), centered), Paragraph('19. října 2026', centered),
                  Paragraph('Festival: 1.–3. srpna 2027', centered), PageBreak(),
                  Paragraph('Obsah', h2)])
    for kind, value in items:
        if kind == 'h2':
            story.append(Paragraph(escape(value), style))
    story.extend([Spacer(1, 16), Paragraph('Povinným zdrojovým dokumentem je documentation.md. Tato verze zachovává stejný obsah v přenosném formátu pro čtení a odevzdání.', style), PageBreak()])
    for kind, value in items:
        if kind in ('h2', 'h3'):
            story.append(Paragraph(escape(value), h2 if kind == 'h2' else h3))
        elif kind == 'table':
            rows = [[Paragraph(('<b>' if i == 0 else '') + pdf_inline(text) + ('</b>' if i == 0 else ''), cell) for text in row] for i, row in enumerate(value)]
            table = Table(rows, colWidths=[width/len(value[0])]*len(value[0]), repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F3DFD7')),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F5F5F5')]),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('LEFTPADDING', (0,0), (-1,-1), 7), ('RIGHTPADDING', (0,0), (-1,-1), 7),
                ('TOPPADDING', (0,0), (-1,-1), 7), ('BOTTOMPADDING', (0,0), (-1,-1), 7),
                ('LINEBELOW', (0,0), (-1,0), .7, colors.HexColor('#'+ORANGE)),
            ]))
            story.extend([table, Spacer(1, 10)])
        elif kind == 'image':
            caption, path = value
            if path.exists():
                image_width, image_height = ImageReader(str(path)).getSize()
                scale = min(width/image_width, 560/image_height)
                image = Image(str(path), width=image_width*scale, height=image_height*scale)
                story.append(KeepTogether([Spacer(1, 8), image, Spacer(1, 6), Paragraph(escape(caption), caption_style)]))
        elif kind == 'code':
            story.append(Paragraph('<br/>'.join(escape(line).replace(' ', '&#160;') for line in value.splitlines()), ParagraphStyle('code', parent=style, fontSize=9, leading=13, leftIndent=8, spaceAfter=18, backColor=colors.HexColor('#F5F5F5'), borderPadding=8)))
        else:
            prefix = '• ' if kind == 'bullet' else ''
            story.append(Paragraph(prefix + pdf_inline(value), style))

    def page_chrome(canvas, document):
        if document.page == 1:
            return
        canvas.saveState()
        canvas.setFont('PulseDoc', 8)
        canvas.setFillColor(colors.HexColor('#'+MUTED))
        canvas.drawString(60, 813, 'PRAGUE PULSE FEST  /  WEBDESIGN')
        canvas.drawRightString(535, 28, f'Dokumentace projektu · {document.page}')
        canvas.setStrokeColor(colors.HexColor('#'+ORANGE))
        canvas.line(60, 803, 535, 803)
        canvas.restoreState()

    pdf = SimpleDocTemplate(str(output), pagesize=(595.28, 841.89), leftMargin=60,
                            rightMargin=60, topMargin=52, bottomMargin=48,
                            title=TITLE+' — dokumentace projektu Webdesign', author=AUTHORS)
    pdf.build(story, onFirstPage=page_chrome, onLaterPages=page_chrome)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=ROOT / 'documentation.md')
    parser.add_argument('--output', type=Path, default=ROOT / 'docs')
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    items = list(blocks(args.source.read_text(encoding='utf-8')))
    render_docx(items, args.output / 'documentation.docx')
    render_pdf(items, args.output / 'documentation.pdf')
    print(f'Dokumentace: {args.output / "documentation.docx"}')
    print(f'Dokumentace: {args.output / "documentation.pdf"}')


if __name__ == '__main__':
    main()
