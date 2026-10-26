"""Validate generated Moodle artifacts and actual prototype scenarios (development only)."""
from pathlib import Path
from threading import Thread
from functools import partial
from http.server import SimpleHTTPRequestHandler,ThreadingHTTPServer
import argparse,hashlib,json,os,zipfile,re
import fitz
from pptx import Presentation
from playwright.sync_api import sync_playwright
parser=argparse.ArgumentParser();parser.add_argument('--folder',type=Path,default=Path('odevzdani'));args=parser.parse_args();root=args.folder.resolve()
checks=[]
def check(ok,label):
 assert ok,label
 checks.append(label);print('PASS',label)
project=Path(__file__).resolve().parents[1]
manifest=json.loads((root/'manifest.json').read_text(encoding='utf-8'))
check(all((root/item['path']).is_file()
          and (root/item['path']).stat().st_size==item['bytes']
          and hashlib.sha256((root/item['path']).read_bytes()).hexdigest()==item['sha256']
          for item in manifest['files']), 'Submission manifest matches every generated artifact')
with zipfile.ZipFile(root/'cviceni-02-k-odevzdani.zip') as archive:
 check(archive.testzip() is None,'C02 upload ZIP has no damaged entries')
 check(set(archive.namelist())=={'project-specification.md','data.xml','festival.xsd','prezentace.pptx'},
       'C02 upload ZIP contains exactly the four intended files, without duplicate presentation or notes')
 check(all(archive.read(name)==(root/'cviceni-02-specifikace'/name).read_bytes() for name in archive.namelist()),
       'C02 upload ZIP agrees with the current specification, data, schema and presentation')
check({path.name for path in (root/'cviceni-02-specifikace').iterdir()}=={
 'project-specification.md','data.xml','festival.xsd','prezentace.pptx'},
 'C02 upload directory contains only the four files intended for Moodle')
prototype_files={'index.html','data.js','prototype.css','prototype.js'}
c03=root/'cviceni-03-prototyp'
check({path.name for path in c03.iterdir()}=={'progress-02.md','prototype'},
      'C03 upload directory contains only the progress report and prototype')
check({path.name for path in (c03/'prototype').iterdir()}==prototype_files,
      'C03 prototype contains its four complete assets without a README or presentation')
check((c03/'prototype/data.js').read_bytes()==(project/'web/js/data.js').read_bytes(),
      'C03 prototype uses the same current festival data as the finished website')
with zipfile.ZipFile(root/'cviceni-03-k-odevzdani.zip') as archive:
 check(archive.testzip() is None,'C03 upload ZIP has no damaged entries')
 check(set(archive.namelist())=={'progress-02.md',*(f'prototype/{name}' for name in prototype_files)},
       'C03 upload ZIP contains exactly the report and four prototype assets')
 check(all(archive.read(name)==(c03/name).read_bytes() for name in archive.namelist()),
       'C03 upload ZIP agrees with the tested upload directory')
for relative in ['cviceni-02-specifikace','cviceni-05-review-ready','finalni-projekt/projekt']:
 folder=root/relative
 for name in ['project-specification.md','data.xml','festival.xsd']:
  check((folder/name).read_bytes()==(project/name).read_bytes(),relative+': '+name+' agrees with the current source')
if manifest['complete_web_copies']:
 expected={path.relative_to(project/'web'):path.read_bytes() for path in (project/'web').rglob('*') if path.is_file()}
 for relative in ['cviceni-04-implementace/web','cviceni-05-review-ready/web','finalni-projekt/projekt/web']:
  delivered={path.relative_to(root/relative):path.read_bytes() for path in (root/relative).rglob('*') if path.is_file()}
  check(delivered==expected,relative+': complete website agrees with the current tested source')
for folder in sorted(path for path in root.glob('cviceni-*') if path.is_dir()):
 number=folder.name.split('-')[1]
 separate=number in {'02','03'}
 pdf_path=root/f'nahledy/cviceni-{number}.pdf' if separate else folder/'prezentace.pdf'
 notes_path=root/f'poznamky/cviceni-{number}.md' if separate else folder/'presenter-notes.md'
 ppt_path=root/'prezentace/cviceni-03.pptx' if number=='03' else folder/'prezentace.pptx'
 pdf=fitz.open(pdf_path);ppt=Presentation(ppt_path)
 check(len(pdf)==6 and len(ppt.slides)==6,folder.name+': six PDF pages and editable PPTX slides')
 for i,(page,slide) in enumerate(zip(pdf,ppt.slides),1):
  check(bool(page.get_text().strip()) and bool(slide.notes_slide.notes_text_frame.text.strip()),folder.name+f': slide{i} text and presenter notes')
  # Inspect the produced files: neither export may silently clip slide text or
  # lose a paragraph while the other format still appears correct.
  spans=[span for block in page.get_text('dict')['blocks'] if 'lines' in block for line in block['lines'] for span in line['spans']]
  check(all(-.5<=s['bbox'][0] and -.5<=s['bbox'][1] and s['bbox'][2]<=page.rect.width+.5 and s['bbox'][3]<=page.rect.height+.5 for s in spans),folder.name+f': slide{i} PDF text stays inside page')
  ppt_text=' '.join(shape.text for shape in slide.shapes if shape.has_text_frame)
  compact=lambda text:re.sub(r'\s+','',text)
  check(compact(page.get_text())==compact(ppt_text),folder.name+f': slide{i} PDF/PPTX visible text agrees')
 check('PULSE' in pdf[0].get_text() and '2027' in pdf[0].get_text(),folder.name+': Prague Pulse 2027 identity')
 check(all(name in pdf[0].get_text() and name in ppt.core_properties.author for name in ['Michal Ladra','Matěj Mlejnek','Jan Martínek']),folder.name+': all three authors in cover and metadata')
 check(notes_path.is_file(),folder.name+': presenter notes kept outside minimal upload directories')
 with zipfile.ZipFile(ppt_path) as archive:check(archive.testzip() is None,folder.name+': PowerPoint archive integrity')
 pdf.close()
class Quiet(SimpleHTTPRequestHandler):
 def log_message(self,*args):pass
server=ThreadingHTTPServer(('127.0.0.1',0),partial(Quiet,directory=str(root)));Thread(target=server.serve_forever,daemon=True).start()
base=f'http://127.0.0.1:{server.server_port}/cviceni-03-prototyp/prototype/index.html'
file_result='Neprovedeno.'
try:
 with sync_playwright() as p:
  # Prefer an explicit browser or the cloud image; use Playwright's bundled
  # Chromium on computers without /usr/bin/chromium, including Windows/macOS.
  chromium_path=os.environ.get('CHROMIUM_PATH')
  if not chromium_path and Path('/usr/bin/chromium').is_file():chromium_path='/usr/bin/chromium'
  launch_options={'args':['--no-sandbox']}
  if chromium_path:launch_options['executable_path']=chromium_path
  browser=p.chromium.launch(**launch_options);page=browser.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
  for width in [320,375,768,1024,1440]:
   page.set_viewport_size({'width':width,'height':900});page.goto(base)
   for view in ['home','program','performers','venues','info','tickets','cart']:
    page.goto(base+'#'+view)
    expected_title={'home':'Úvod','program':'Program','performers':'Účinkující','venues':'Scény','info':'Informace','tickets':'Vstupenky','cart':'Košík'}[view]
    page.wait_for_function('(title)=>document.title.startsWith(title)',arg=expected_title)
    check(page.evaluate('document.documentElement.scrollWidth<=innerWidth'),f'Prototype fits{width}px/{view}')
  page.goto(base+'#program')
  page.wait_for_selector('#program-list article')
  check(page.locator('#program-list article').count()==60,'Prototype loads60program items')
  events=page.evaluate('FESTIVAL_DATA.events')
  actual=page.locator('#program-list article').evaluate_all('(nodes)=>nodes.map(node=>node.dataset.eventId)')
  expected=[event['id'] for event in sorted(events,key=lambda event:(event['date'],event['startTime'],event['venueName']))]
  check(actual==expected,'Prototype orders the full programme by date, time and stage')
  old_hash=page.evaluate('location.hash')
  page.locator('.skip').focus();page.keyboard.press('Enter')
  check(page.evaluate('location.hash')==old_hash and page.evaluate('document.activeElement.id')=='main'
        and page.locator('#program-list article').count()==60,'Prototype skip link focuses content without changing the routed view')
  page.locator('#day').select_option('d1');page.locator('#venue').select_option('h2')
  expected_ids={event['id'] for event in events if event['dayId']=='d1' and event['venueId']=='h2'}
  check(set(page.locator('#program-list article').evaluate_all('(nodes)=>nodes.map(node=>node.dataset.eventId)'))==expected_ids
        and len(expected_ids)==7,'Prototype S1 combines the exact first-day River Stage data')
  saved_id=page.locator('[data-event]').first.get_attribute('data-event')
  page.locator('[data-event]').first.click();page.locator('#only-saved').check()
  check(page.locator('#program-list article').count()==1,'Prototype scenarioS1 saved program')
  page.reload();page.locator('#only-saved').check()
  check(page.locator('#program-list article').count()==1
        and page.locator('[data-event]').get_attribute('data-event')==saved_id,'Prototype S1 preserves the exact saved concert after reload')
  page.locator('[data-event]').click()
  check(page.locator('#program-list article').count()==0 and page.evaluate('document.activeElement.id')=='only-saved',
        'Removing the last saved concert leaves useful keyboard focus')
  page.goto(base+'#venues')
  page.locator('main article').filter(has=page.get_by_role('heading',name='River Stage',exact=True)).locator('a').click()
  check(page.locator('#venue').input_value()=='h2'
        and all(event_id in {event['id'] for event in events if event['venueId']=='h2'}
                for event_id in page.locator('#program-list article').evaluate_all('(nodes)=>nodes.map(node=>node.dataset.eventId)')),
        'A stage card opens its own programme with the stage filter selected')
  page.goto(base+'#program?venue=h3&day=d3')
  page.wait_for_function("document.querySelector('#venue')?.value==='h3' && document.querySelector('#day')?.value==='d3'")
  check(page.locator('#venue').input_value()=='h3' and page.locator('#day').input_value()=='d3'
        and page.locator('#program-list article').count()==len([event for event in events if event['venueId']=='h3' and event['dayId']=='d3']),
        'Prototype links can preselect both the stage and day')
  page.goto(base+'#performers');page.wait_for_selector('article[data-artist-id]')
  check(page.locator('main article').count()==60,'Prototype loads60artists')
  page.locator('#query').fill('darkz');page.locator('#search-form button').click()
  page.wait_for_function("document.title.startsWith('Vyhledávání')")
  artist=page.locator('main article').filter(has=page.get_by_role('heading',name='DARKZY',exact=True))
  artist.locator('[data-artist]').click()
  check(page.locator('dialog').is_visible() and page.locator('#dialog-title').inner_text()=='DARKZY',
        'Prototype S2 finds the intended artist by a case-insensitive partial name')
  artist_id=page.evaluate('FESTIVAL_DATA.performers.find(artist=>artist.name==="DARKZY").id')
  related=[event for event in events if artist_id in event['performerIds']]
  check(all(event['venueName'] in page.locator('#dialog-content').inner_text()
            and event['startTime'][:5] in page.locator('#dialog-content').inner_text() for event in related),
        'Prototype S2 detail displays the canonical show time and stage')
  page.locator('dialog a[href^="#program?artist="]').click();page.wait_for_selector('#program-list article')
  check(set(page.locator('#program-list article').evaluate_all('(nodes)=>nodes.map(node=>node.dataset.eventId)'))=={event['id'] for event in related},
        'Prototype S2 opens only the shows linked to the selected artist ID')
  page.locator('#query').fill('DARKZY');page.locator('#search-form button').click()
  page.wait_for_function("document.title.startsWith('Vyhledávání')")
  page.locator('main article').filter(has=page.get_by_role('heading',name='DARKZY',exact=True)).locator('[data-artist]').click()
  page.locator('dialog a[href^="#tickets?event="]').first.click()
  page.wait_for_selector('#prototype-add-cart')
  show=related[0]
  check(page.locator('#ticket-day').input_value()==show['dayId'] and page.locator('#ticket-venue').input_value()==show['venueId'],
        'Prototype S2 carries the actual show day and stage into the ticket selection')
  catalogue=page.evaluate('FESTIVAL_DATA.venues')
  venue=next(venue for venue in catalogue if venue['id']==show['venueId'])
  unit=venue['ticketing']['prices']['standing']['day']
  compact=lambda value:re.sub(r'\s+','',value)
  check(compact(f'{unit} Kč') in compact(page.locator('#ticket-total').inner_text()),
        'Prototype ticket price comes from the selected stage catalogue')
  page.locator('#ticket-quantity').fill('2')
  check(compact(f'{unit*2} Kč') in compact(page.locator('#ticket-total').inner_text()),
        'Prototype ticket quantity immediately recalculates its price')
  page.locator('#prototype-add-cart').click();page.wait_for_selector('#cart-total')
  check(compact(f'{unit*2} Kč') in compact(page.locator('#cart-total').inner_text())
        and page.locator('#cart-count').inner_text()=='2' and show['venueName'] in page.locator('main').inner_text(),
        'Prototype S2 finishes in a cart with the selected stage, quantity and price')
  page.reload();page.wait_for_selector('#cart-total')
  check(page.locator('[data-cart-item]').count()==1 and page.locator('#cart-count').inner_text()=='2',
        'Prototype cart preserves the selected ticket after reload')
  page.set_viewport_size({'width':320,'height':900})
  check(page.evaluate('document.documentElement.scrollWidth<=innerWidth'),'Filled prototype cart fits320px')
  page.locator('main a[href="#tickets"]').click();page.wait_for_selector('#prototype-add-cart')
  page.locator('#ticket-venue').select_option(show['venueId']);page.locator('#ticket-day').select_option(show['dayId'])
  check(int(page.locator('#ticket-remaining').get_attribute('data-remaining'))==venue['ticketing']['standingCapacity']-2,
        'Prototype availability accounts for the existing local selection on the same stage and day')
  page.locator('nav a[href="#cart"]').click();page.wait_for_selector('[data-remove-cart]')
  page.locator('[data-remove-cart]').click()
  check(page.locator('[data-cart-item]').count()==0 and page.locator('#cart-count').inner_text()=='0'
        and 'Košík je prázdný' in page.locator('main').inner_text(),
        'Removing a prototype cart item updates the total and empty-cart state')
  page.locator('#query').fill('<script>alert(1)</script>');page.locator('#search-form button').click()
  page.wait_for_function("document.querySelector('main').textContent.includes('Nic jsme')")
  check('Nic jsme' in page.locator('main').inner_text() and not page.locator('main script').count(),
        'Prototype search treats markup as text')
  page.goto(base+'#program');page.wait_for_function("document.title.startsWith('Program')")
  page.locator('#query').fill('');page.locator('#search-form button').click()
  page.wait_for_function("document.title.startsWith('Vyhledávání')")
  check(page.title().startswith('Vyhledávání'),'Empty search updates page metadata')
  check(not errors,'Prototype has no runtime errors')
  file_context=browser.new_context();file_page=file_context.new_page()
  try:
   file_page.goto((c03/'prototype/index.html').as_uri())
   blocked=file_page.evaluate('location.href').startswith('chrome-error:')
   if blocked:file_result='Přímé file:// otevření přeskočeno kvůli politice místního prohlížeče.'
   else:
    check(file_page.locator('main h1').inner_text()=='Prague Pulse Fest','Prototype starts directly from disk without a server')
    file_page.locator('nav a[href="#program"]').click()
    file_page.wait_for_selector('#program-list article')
    check(file_page.locator('#program-list article').count()==60,'Direct file prototype loads its complete offline programme')
    file_result='Přímé file:// otevření a offline program prošly.'
  except Exception as error:
   if 'ERR_BLOCKED_BY_ADMINISTRATOR' not in str(error):raise
   file_result='Přímé file:// otevření přeskočeno kvůli politice místního prohlížeče.'
  file_context.close();browser.close()
finally:server.shutdown();server.server_close()
(root/'prototype-validation.md').write_text('# Ověření podkladů a prototypu\n\n'+str(len(checks))+' kontrol prošlo.\n\n'+'\n'.join('- '+c for c in checks)+'\n\nKontroly proběhly v Chromium přes vývojový statický server. '+file_result+'\n',encoding='utf-8')
print(f'{len(checks)} course/prototype checks passed.')
